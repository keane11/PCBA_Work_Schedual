"""Build small, readable Excel exports without evaluating user supplied text."""

from __future__ import annotations

from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# Column tuple: row key, Chinese heading, Excel value kind, approximate width.
Column = tuple[str, str, str, int]


def _write_value(cell: Any, value: Any, kind: str) -> None:
    if value is None or value == "":
        return
    if kind == "date":
        try:
            cell.value = date.fromisoformat(str(value)[:10])
            cell.number_format = "yyyy-mm-dd"
            return
        except ValueError:
            pass
    elif kind == "number":
        try:
            amount = Decimal(str(value))
            if amount.is_finite():
                cell.value = float(amount)
                cell.number_format = "#,##0.########"
                return
        except (InvalidOperation, ValueError, TypeError):
            pass
    # Force all other values to text. openpyxl otherwise treats leading '=' as a
    # formula, and Excel otherwise drops leading zeroes from order/item codes.
    cell.value = str(value)
    cell.data_type = "s"
    cell.number_format = "@"


def build_workbook(sheets: list[tuple[str, list[Column], list[dict]]], notes: list[tuple[str, Any]]) -> bytes:
    workbook = Workbook()
    note_sheet = workbook.active
    note_sheet.title = "导出说明"
    note_sheet.column_dimensions["A"].width = 24
    note_sheet.column_dimensions["B"].width = 72
    for row_number, (label, value) in enumerate(notes, 1):
        _write_value(note_sheet.cell(row_number, 1), label, "text")
        _write_value(note_sheet.cell(row_number, 2), value, "text")
        note_sheet.cell(row_number, 1).font = Font(bold=True, color="155C75")
        note_sheet.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")
    note_sheet.freeze_panes = "A2"

    for title, columns, records in sheets:
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = "A2"
        for col_number, (_, heading, _, width) in enumerate(columns, 1):
            cell = sheet.cell(1, col_number)
            _write_value(cell, heading, "text")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="155C75")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(col_number)].width = width
        sheet.row_dimensions[1].height = 30
        for row_number, record in enumerate(records, 2):
            for col_number, (key, _, kind, _) in enumerate(columns, 1):
                cell = sheet.cell(row_number, col_number)
                _write_value(cell, record.get(key), kind)
                if row_number % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F3F7F8")
                cell.alignment = Alignment(vertical="top")
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(records) + 1)}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

"""Excel exports keep project stages and customer delivery batches distinct."""

from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook

from server import app as backend


def values(sheet: object) -> list[dict]:
    rows = list(sheet.iter_rows(values_only=True))
    return [dict(zip(rows[0], row)) for row in rows[1:]]


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.previous_path = backend.DB_PATH
        backend.DB_PATH = Path(self.temp.name) / "pcba.db"
        backend.init_db()
        now = backend.utcnow()
        with backend.writing() as conn:
            user = conn.execute("""INSERT INTO users(username,display_name,role,password_salt,password_hash,created_at)
                                   VALUES('admin','管理员','admin','salt','hash',?)""", (now,)).lastrowid
            customer = conn.execute("""INSERT INTO orders(order_no,customer,project_name,project_mode,owner_id,created_at,updated_at)
                                       VALUES('000123','=试验客户','项目 A','customer',?,?,?)""", (user, now, now)).lastrowid
            line = conn.execute("""INSERT INTO order_lines(order_id,line_no,product_code,product_name,quantity,created_at,updated_at)
                                   VALUES(?,'01','+001A','控制板','10',?,?)""", (customer, now, now)).lastrowid
            shared = conn.execute("""INSERT INTO production_batches(order_line_id,batch_no,quantity,phase,vendor,
                                    commit_start,commit_finish,schedule_start,created_at,updated_at)
                                    VALUES(?,'0001','10','pp','外协甲','2026-10-01','2026-10-08','2026-09-30',?,?)""",
                                  (line, now, now)).lastrowid
            conn.execute("""INSERT INTO stages(production_batch_id,sequence,name,planned_start,planned_finish,
                          actual_finish,status,owner_name,completed_by,created_at,updated_at)
                          VALUES(?,1,'SMT','2026-10-01','2026-10-03','2026-10-04','done','唐林林',?,?,?)""",
                         (shared, user, now, now))
            for number, quantity in (("D01", "6"), ("D02", "4")):
                delivery = conn.execute("""INSERT INTO delivery_batches(order_line_id,batch_no,quantity,due_date,promised_date,
                                         created_at,updated_at) VALUES(?,?,?,'2026-10-15','2026-10-13',?,?)""",
                                        (line, number, quantity, now, now)).lastrowid
                conn.execute("""INSERT INTO allocations(production_batch_id,delivery_batch_id,quantity,created_at)
                              VALUES(?,?,?,?)""", (shared, delivery, quantity, now))
            internal = conn.execute("""INSERT INTO orders(order_no,customer,project_name,project_mode,created_at,updated_at)
                                       VALUES('PCBA-PROJ-1','内部研发','研发项目','internal',?,?)""", (now, now)).lastrowid
            self.internal_line = conn.execute("""INSERT INTO order_lines(order_id,line_no,product_code,product_name,quantity,
                                               board_schedule_mode,created_at,updated_at)
                                               VALUES(?,'1','00045','样板','10','independent',?,?)""",
                                              (internal, now, now)).lastrowid
            independent = conn.execute("""INSERT INTO production_batches(order_line_id,batch_no,quantity,phase,vendor,
                                          board_schedule_mode,bom_scope,schedule_start,created_at,updated_at)
                                          VALUES(?,'RD01','10','rd','批次厂商','independent','multi','2026-10-01',?,?)""",
                                       (self.internal_line, now, now)).lastrowid
            board_definition = conn.execute("""INSERT INTO order_line_boards(order_line_id,board_ref,board_name,
                                              bom_product_code,quantity_per_product,created_at,updated_at)
                                              VALUES(?,'001','主板','000PCB','1',?,?)""",
                                            (self.internal_line, now, now)).lastrowid
            board = conn.execute("""INSERT INTO batch_boards(production_batch_id,order_line_board_id,board_ref,board_name,
                                    bom_product_code,quantity_per_product,planned_quantity,created_at,updated_at)
                                    VALUES(?,?,'001','主板','000PCB','1','10',?,?)""",
                                 (independent, board_definition, now, now)).lastrowid
            conn.execute("""INSERT INTO board_stages(batch_board_id,sequence,name,planned_start,planned_finish,
                          actual_finish,status,owner_name,completed_by,created_at,updated_at)
                          VALUES(?,1,'样板到料','2026-10-02','2026-10-05','2026-10-06','done','白雪飞',?,?,?)""",
                         (board, user, now, now))
            conn.execute("""INSERT INTO events(entity_type,entity_id,kind,event_date,quantity,actor_id,created_at)
                          VALUES('batch_board',?,'finished','2026-10-06','2',?,?)""", (board, user, now))

    def tearDown(self):
        backend.DB_PATH = self.previous_path
        self.temp.cleanup()

    def workbook(self, response):
        self.assertEqual(response.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment", response.headers["Content-Disposition"])
        return load_workbook(BytesIO(response.body))

    def test_project_export_includes_shared_and_independent_nodes(self):
        workbook = self.workbook(backend.export_project_schedule())
        self.assertEqual(workbook.sheetnames, ["导出说明", "项目清单", "生产批次", "板件排期", "生产节点"])
        self.assertEqual(len(values(workbook["项目清单"])), 2)
        nodes = values(workbook["生产节点"])
        self.assertEqual({row["生产节点"] for row in nodes}, {"SMT", "样板到料"})
        self.assertEqual({row["节点责任人"] for row in nodes}, {"唐林林", "白雪飞"})
        board = values(workbook["板件排期"])[0]
        self.assertEqual(board["板卡编码"], "000PCB")
        self.assertEqual(board["板卡完工数量"], 2)
        self.assertIsNone(board["外协厂"])  # Independent board cannot inherit a batch promise.
        self.assertIsNone(board["承诺开工"])
        self.assertEqual(board["计划完成"].date().isoformat(), "2026-10-05")
        self.assertEqual(workbook["项目清单"].freeze_panes, "A2")
        self.assertIsNotNone(workbook["项目清单"].auto_filter.ref)

        selected = self.workbook(backend.export_project_schedule(order_line_id=self.internal_line, phase="rd"))
        self.assertEqual(len(values(selected["项目清单"])), 1)
        self.assertEqual(values(selected["生产节点"])[0]["板卡标识"], "001")
        self.assertEqual(values(selected["生产节点"])[0]["完成状态"], "已完成")

    def test_customer_export_preserves_identifiers_and_split_deliveries(self):
        workbook = self.workbook(backend.export_customer_orders(q="+001A"))
        self.assertEqual(workbook.sheetnames, ["导出说明", "客户订单", "交付批次", "生产交付关联"])
        self.assertEqual(len(values(workbook["客户订单"])), 1)  # Internal R&D is excluded.
        deliveries = values(workbook["交付批次"])
        self.assertEqual([row["交付批次"] for row in deliveries], ["D01", "D02"])
        self.assertEqual(sum(row["本批交付数量"] for row in deliveries), 10)
        self.assertEqual([row["本批已交付数量"] for row in deliveries], [0, 0])
        self.assertTrue(all(row["正式预测交付日"] is None and row["待确认原因"] for row in deliveries))
        self.assertEqual(len(values(workbook["生产交付关联"])), 2)
        order_sheet = workbook["客户订单"]
        headings = {cell.value: cell.column for cell in order_sheet[1]}
        self.assertEqual(order_sheet.cell(2, headings["客户订单号"]).value, "000123")
        self.assertEqual(order_sheet.cell(2, headings["客户订单号"]).data_type, "s")
        self.assertEqual(order_sheet.cell(2, headings["客户"]).value, "=试验客户")
        self.assertEqual(order_sheet.cell(2, headings["客户"]).data_type, "s")
        self.assertEqual(order_sheet.cell(2, headings["产品编码"]).value, "+001A")
        self.assertEqual(order_sheet.cell(2, headings["产品编码"]).data_type, "s")
        self.assertEqual(workbook["交付批次"].cell(2, 8).number_format, "#,##0.########")

    def test_export_filters_report_empty_or_invalid_range(self):
        for call, status in ((lambda: backend.export_project_schedule(phase="bad"), 422),
                             (lambda: backend.export_project_schedule(order_line_id=9999), 404),
                             (lambda: backend.export_customer_orders(q="not-found"), 404)):
            with self.assertRaises(HTTPException) as caught:
                call()
            self.assertEqual(caught.exception.status_code, status)
